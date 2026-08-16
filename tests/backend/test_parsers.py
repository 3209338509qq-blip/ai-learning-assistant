"""解析器测试：TXT / MD / DOCX / PDF。"""
from pathlib import Path

from app.services.parsers import (
    parse_docx,
    parse_markdown,
    parse_pdf,
    parse_txt,
)


def test_parse_txt_single_section(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("第一段内容。\n\n第二段内容。", encoding="utf-8")
    doc = parse_txt(str(p))
    assert doc.file_type == "txt"
    assert len(doc.sections) == 1
    assert "第一段内容" in doc.full_text


def test_parse_markdown_chapters(tmp_path):
    p = tmp_path / "notes.md"
    p.write_text(
        "# Python 基础\n\n介绍。\n\n## 装饰器\n\n装饰器是一种函数。\n\n## 生成器\n\nyield 关键字。",
        encoding="utf-8",
    )
    doc = parse_markdown(str(p))
    paths = [s.path for s in doc.sections]
    assert "Python 基础 / 装饰器" in paths
    assert "Python 基础 / 生成器" in paths
    deco = [s for s in doc.sections if s.path == "Python 基础 / 装饰器"][0]
    assert "装饰器是一种函数" in deco.text


def test_parse_docx_headings(tmp_path):
    import docx

    p = tmp_path / "a.docx"
    d = docx.Document()
    d.add_heading("第一章 引言", level=1)
    d.add_paragraph("这是第一章正文。")
    d.add_heading("1.1 背景", level=2)
    d.add_paragraph("背景内容。")
    d.save(str(p))

    doc = parse_docx(str(p))
    paths = [s.path for s in doc.sections]
    assert "第一章 引言 / 1.1 背景" in paths
    assert any("这是第一章正文" in s.text for s in doc.sections)


def test_parse_pdf_with_pages(tmp_path):
    import fitz

    p = tmp_path / "a.pdf"
    pdf = fitz.open()
    page1 = pdf.new_page()
    page1.insert_text((72, 72), "Chapter 1 Introduction", fontsize=18)
    page1.insert_text((72, 100), "This is page one content.")
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "More content on page two.")
    pdf.save(str(p))
    pdf.close()

    doc = parse_pdf(str(p))
    assert doc.page_count == 2
    # 大字号标题行被识别为章节
    assert any("Chapter 1" in s.title for s in doc.sections)
    # 正文带页码
    s = [s for s in doc.sections if "page one content" in s.text]
    assert s and s[0].page == 1
    assert any("More content on page two" in s.text for s in doc.sections)


def test_unsupported_type_raises(tmp_path):
    import pytest

    from app.services.parsers import parse_document

    p = tmp_path / "a.xyz"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_document(str(p), "xyz")
