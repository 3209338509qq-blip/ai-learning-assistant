"""文档解析：PDF / DOCX / Markdown / TXT → 带章节与页码的结构化文本。

统一输出 ParsedDocument（sections 列表），每个 section 记录：
- title / level / path（章节路径，用于溯源）
- page（起始页码，PDF 有效；DOCX/MD/TXT 无页码概念记 0/1）
- text（章节正文）
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

SUPPORTED_TYPES = {"pdf", "docx", "md", "txt"}

# 常见章节编号（用于 PDF 无样式信息时的标题识别）
_HEADING_NUM = re.compile(
    r"^(第[一二三四五六七八九十百零0-9]+[章节篇卷]|[0-9]+(?:\.[0-9]+){0,3}[ 　\t]|"
    r"Chapter\s+[0-9IVX]+|附录[ABC]?[ 　:]|[A-Z]\.[0-9]+[ 　])"
)
_CHAPTER_NUM = re.compile(r"^第[一二三四五六七八九十百零0-9]+[章节篇卷]")


@dataclass
class ParsedSection:
    title: str
    level: int
    path: str
    page: int
    text: str = ""
    order: int = 0


@dataclass
class ParsedDocument:
    file_type: str
    page_count: int
    sections: list[ParsedSection] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n".join(s.text for s in self.sections)


def _heading_level(title: str) -> int:
    """从标题编号推断层级：'第X章'→1，'1.2'→2，'1.2.3'→3；无编号→1。"""
    m = re.match(r"^(\d+(?:\.\d+)*)", title.strip())
    if m:
        return m.group(1).count(".") + 1
    if _CHAPTER_NUM.match(title.strip()):
        return 1
    return 1


class SectionBuilder:
    """按标题流构建章节树：标题行开新章节，正文归入当前章节。"""

    def __init__(self, default_title: str = "全文"):
        self.sections: list[ParsedSection] = []
        self._stack: list[tuple[str, int]] = []
        self._current: ParsedSection | None = None
        self._default_title = default_title

    def add_heading(self, title: str, level: int, page: int, order: int) -> None:
        title = title.strip()[:200]
        if not title:
            return
        while self._stack and self._stack[-1][1] >= level:
            self._stack.pop()
        self._stack.append((title, level))
        path = " / ".join(t for t, _ in self._stack)
        self._current = ParsedSection(title=title, level=level, path=path, page=page, order=order)
        self.sections.append(self._current)

    def add_text(self, text: str, page: int) -> None:
        text = text.strip()
        if not text:
            return
        if self._current is None:
            self.add_heading(self._default_title, 0, page, 0)
        self._current.text += text + "\n"

    def build(self) -> list[ParsedSection]:
        sections = [s for s in self.sections if s.text.strip()]
        if not sections:
            sections = [ParsedSection(self._default_title, 0, self._default_title, 1)]
        for i, s in enumerate(sections):
            s.order = i
        return sections


# ---------------- PDF（PyMuPDF）----------------

def _pdf_line_text(line: dict) -> str:
    """合并 span 文本，依据坐标间隙补空格。"""
    out = ""
    prev_x1: float | None = None
    for span in line["spans"]:
        t = span["text"]
        x0 = span["bbox"][0]
        if prev_x1 is not None and x0 - prev_x1 > 1.5 and not out.endswith(" "):
            out += " "
        out += t
        prev_x1 = span["bbox"][2]
    return out.strip()


def _looks_like_heading(text: str, size: float, avg: float, bold: bool) -> bool:
    text = text.strip()
    if not text or len(text) > 90:
        return False
    if _HEADING_NUM.match(text):
        return True
    # 字号明显大于正文 或 加粗 且 行较短
    if (size >= avg * 1.15 or bold) and len(text) <= 50:
        return True
    return False


def parse_pdf(path: str) -> ParsedDocument:
    import fitz  # PyMuPDF

    pdf = fitz.open(path)
    builder = SectionBuilder()
    order = 0

    sizes: list[float] = []
    for page in pdf:
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block["lines"]:
                for span in line["spans"]:
                    if span["text"].strip():
                        sizes.append(span["size"])
    avg = sum(sizes) / len(sizes) if sizes else 12.0

    for page in pdf:
        page_no = page.number + 1
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block["lines"]:
                spans = line["spans"]
                if not spans:
                    continue
                text = _pdf_line_text(line)
                if not text:
                    continue
                size = max(s["size"] for s in spans)
                bold = any(s["flags"] & 16 for s in spans)
                if _looks_like_heading(text, size, avg, bold):
                    order += 1
                    builder.add_heading(text, _heading_level(text), page_no, order)
                else:
                    builder.add_text(text, page_no)

    page_count = pdf.page_count
    pdf.close()
    return ParsedDocument(file_type="pdf", page_count=page_count, sections=builder.build())


# ---------------- DOCX（python-docx）----------------

def parse_docx(path: str) -> ParsedDocument:
    import docx

    d = docx.Document(path)
    builder = SectionBuilder()
    order = 0

    def style_level(style_name: str | None) -> int | None:
        if not style_name:
            return None
        name = style_name.lower()
        if name.startswith("heading"):
            digits = re.sub(r"\D", "", name)
            return int(digits) if digits else 1
        if name in ("title", "标题"):
            return 1
        return None

    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        lvl = style_level(para.style.name if para.style else None)
        if lvl is not None:
            order += 1
            builder.add_heading(text, lvl, 0, order)
        else:
            builder.add_text(text, 0)

    return ParsedDocument(file_type="docx", page_count=0, sections=builder.build())


# ---------------- Markdown / TXT ----------------

def _split_markdown(text: str, file_type: str, page_count: int) -> ParsedDocument:
    builder = SectionBuilder()
    order = 0
    for line in text.splitlines():
        stripped = line.strip()
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            order += 1
            builder.add_heading(m.group(2), len(m.group(1)), 0, order)
        else:
            builder.add_text(line, 0)
    return ParsedDocument(file_type=file_type, page_count=page_count, sections=builder.build())


def parse_markdown(path: str) -> ParsedDocument:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return _split_markdown(text, "md", 0)


def parse_txt(path: str) -> ParsedDocument:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return _split_markdown(text, "txt", 0)


# ---------------- 统一入口 ----------------

def parse_document(path: str, file_type: str) -> ParsedDocument:
    if file_type not in SUPPORTED_TYPES:
        raise ValueError(f"不支持的文件类型: {file_type}")
    if file_type == "pdf":
        return parse_pdf(path)
    if file_type == "docx":
        return parse_docx(path)
    if file_type == "md":
        return parse_markdown(path)
    return parse_txt(path)
